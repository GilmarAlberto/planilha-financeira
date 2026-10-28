"""Gera a planilha financeira pessoal em Excel.

A planilha tem cinco abas:
    Como usar      instruções rápidas
    Lançamentos    onde você anota cada receita e cada despesa
    Orçamento      o limite de cada categoria no mês, com alerta verde/amarelo/vermelho
    Resumo mensal  receitas, despesas e saldo de cada mês, com gráfico
    Categorias     a lista de categorias (pode mudar à vontade)

Todas as contas são fórmulas do Excel: depois de gerada, a planilha
funciona sozinha, sem precisar do Python.

Uso:
    python gerar_planilha.py              planilha com lançamentos de exemplo
    python gerar_planilha.py --em-branco  planilha vazia, pronta para usar
    python gerar_planilha.py --substituir gera de novo por cima da que já existe

Por segurança, se a planilha já existe o programa não mexe nela
(ela pode ter os seus lançamentos). Para gerar outra, use --substituir.
"""
import random
import sys
from datetime import date
from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import BarChart, PieChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

ARQUIVO = Path("planilha_financeira.xlsx")
ANO = 2026
LINHAS_DE_LANCAMENTO = 1000          # espaço reservado na aba Lançamentos

# ---------------------------------------------------------------------------
# Categorias e limites do orçamento. Mude aqui para a sua realidade.
# ---------------------------------------------------------------------------
RECEITAS = ["Salário", "Freela", "Outras receitas"]
DESPESAS = {                          # categoria: limite mensal em R$
    "Moradia": 1500,
    "Mercado": 900,
    "Contas (luz, água, internet)": 350,
    "Transporte": 400,
    "Saúde": 300,
    "Educação": 200,
    "Lazer": 300,
    "Outras despesas": 200,
}

# ---------------------------------------------------------------------------
# Cores e estilos (um lugar só, para mudar o visual inteiro de uma vez)
# ---------------------------------------------------------------------------
AZUL_ESCURO = "1F3A5F"
VERDE, VERDE_CLARO = "1D6B35", "DEF7E5"
AMARELO, AMARELO_CLARO = "8A6100", "FFF4CE"
VERMELHO, VERMELHO_CLARO = "A61B1B", "FDE2E2"
CINZA_CLARO = "F4F6F8"
FONTE = "Calibri"
REAIS = '"R$" #,##0.00;[Red]-"R$" #,##0.00'
LINHA_FINA = Side(style="thin", color="D0D7DE")


def preencher(cor):
    return PatternFill("solid", fgColor=cor)


def titulo(aba, texto, subtitulo):
    """Linha 1: título grande. Linha 2: uma frase explicando a aba."""
    aba["A1"] = texto
    aba["A1"].font = Font(name=FONTE, size=16, bold=True, color=AZUL_ESCURO)
    aba["A2"] = subtitulo
    aba["A2"].font = Font(name=FONTE, italic=True, color="666666")
    aba.sheet_view.showGridLines = False


def cabecalho(aba, linha, colunas):
    """colunas = [(nome, largura), ...]. Pinta o cabeçalho de azul."""
    for numero, (nome, largura) in enumerate(colunas, start=1):
        celula = aba.cell(linha, numero, nome)
        celula.font = Font(name=FONTE, bold=True, color="FFFFFF")
        celula.fill = preencher(AZUL_ESCURO)
        celula.alignment = Alignment(horizontal="center", vertical="center")
        aba.column_dimensions[celula.column_letter].width = largura
    aba.row_dimensions[linha].height = 22


def para_impressao(aba, paisagem=True):
    """Ao imprimir (ou salvar em PDF), cabe na largura de uma folha."""
    aba.page_setup.orientation = "landscape" if paisagem else "portrait"
    aba.sheet_properties.pageSetUpPr.fitToPage = True
    aba.page_setup.fitToWidth, aba.page_setup.fitToHeight = 1, 0


def borda(celula):
    celula.border = Border(bottom=LINHA_FINA)


# ---------------------------------------------------------------------------
# Aba Categorias: a lista que alimenta os menus de escolha
# ---------------------------------------------------------------------------
def aba_categorias(livro):
    aba = livro.create_sheet("Categorias")
    titulo(aba, "Categorias", "Mude, apague ou acrescente categorias aqui. Os menus usam esta lista.")
    cabecalho(aba, 4, [("Categoria", 32), ("Tipo", 12)])
    linha = 5
    for nome in RECEITAS:
        aba.cell(linha, 1, nome), aba.cell(linha, 2, "Receita")
        linha += 1
    for nome in DESPESAS:
        aba.cell(linha, 1, nome), aba.cell(linha, 2, "Despesa")
        linha += 1
    for celulas in aba.iter_rows(min_row=5, max_row=linha - 1):
        for celula in celulas:
            celula.font = Font(name=FONTE)
            borda(celula)
    # Deixa 20 linhas a mais no menu, para quem quiser criar categorias novas
    return f"=Categorias!$A$5:$A${linha + 20}"


# ---------------------------------------------------------------------------
# Aba Lançamentos: cada linha é uma receita ou uma despesa
# ---------------------------------------------------------------------------
def aba_lancamentos(livro, lista_de_categorias, exemplos):
    aba = livro.create_sheet("Lançamentos")
    titulo(aba, "Lançamentos", "Anote aqui cada entrada e cada saída de dinheiro. Tipo e Categoria têm menu.")
    cabecalho(aba, 4, [("Data", 13), ("Descrição", 38), ("Tipo", 12),
                       ("Categoria", 30), ("Valor", 15)])
    aba.freeze_panes = "A5"

    ultima = 4 + LINHAS_DE_LANCAMENTO
    for linha in range(5, ultima + 1):
        aba.cell(linha, 1).number_format = "dd/mm/yyyy"
        aba.cell(linha, 5).number_format = REAIS
        for coluna in range(1, 6):
            aba.cell(linha, coluna).font = Font(name=FONTE)
            borda(aba.cell(linha, coluna))
    for linha, lancamento in enumerate(exemplos, start=5):
        for coluna, valor in enumerate(lancamento, start=1):
            aba.cell(linha, coluna, valor)

    # Menus de escolha: o Tipo só aceita Receita ou Despesa; a Categoria vem da aba Categorias
    tipo = DataValidation(type="list", formula1='"Receita,Despesa"', allow_blank=True,
                          error="Escolha Receita ou Despesa.", errorTitle="Tipo inválido")
    categoria = DataValidation(type="list", formula1=lista_de_categorias, allow_blank=True,
                               error="Escolha uma categoria da lista (ou crie na aba Categorias).",
                               errorTitle="Categoria inválida")
    valor = DataValidation(type="decimal", operator="greaterThan", formula1="0", allow_blank=True,
                           error="Digite o valor sem sinal. O Tipo já diz se é entrada ou saída.",
                           errorTitle="Valor inválido")
    data = DataValidation(type="date", operator="greaterThan", formula1="36526", allow_blank=True,
                          error="Digite uma data, por exemplo 05/10/2026.", errorTitle="Data inválida")
    for regra, coluna in ((tipo, "C"), (categoria, "D"), (valor, "E"), (data, "A")):
        aba.add_data_validation(regra)
        regra.add(f"{coluna}5:{coluna}{ultima}")

    # Receita em verde, despesa em vermelho: dá para bater o olho e entender
    faixa = f"A5:E{ultima}"
    aba.conditional_formatting.add(faixa, FormulaRule(
        formula=['$C5="Receita"'], font=Font(color=VERDE)))
    aba.conditional_formatting.add(faixa, FormulaRule(
        formula=['$C5="Despesa"'], font=Font(color=VERMELHO)))
    return ultima


# ---------------------------------------------------------------------------
# Aba Resumo mensal: uma linha por mês, com gráfico de barras
# ---------------------------------------------------------------------------
def soma_do_mes(tipo, inicio_do_mes, ultima):
    """Fórmula do Excel que soma os lançamentos de um tipo dentro de um mês."""
    l = "'Lançamentos'!"
    return (f'=SUMIFS({l}$E$5:$E${ultima},{l}$C$5:$C${ultima},"{tipo}",'
            f'{l}$A$5:$A${ultima},">="&{inicio_do_mes},'
            f'{l}$A$5:$A${ultima},"<"&EDATE({inicio_do_mes},1))')


def aba_resumo(livro, ultima):
    aba = livro.create_sheet("Resumo mensal")
    titulo(aba, "Resumo mensal", "Quanto entrou, quanto saiu e quanto sobrou em cada mês.")
    aba["A3"], aba["B3"] = "Ano:", ANO
    aba["A3"].font = Font(name=FONTE, bold=True)
    aba["B3"].font = Font(name=FONTE, bold=True, color=AZUL_ESCURO)
    aba["B3"].fill = preencher(AMARELO_CLARO)
    aba["C3"] = "← mude o ano aqui"
    aba["C3"].font = Font(name=FONTE, italic=True, color="666666")

    cabecalho(aba, 5, [("Mês", 13), ("Receitas", 15), ("Despesas", 15),
                       ("Saldo do mês", 15), ("Saldo acumulado", 17)])
    for mes in range(1, 13):
        linha = 5 + mes
        aba.cell(linha, 1, f"=DATE($B$3,{mes},1)").number_format = "mmm/yyyy"
        aba.cell(linha, 2, soma_do_mes("Receita", f"$A{linha}", ultima))
        aba.cell(linha, 3, soma_do_mes("Despesa", f"$A{linha}", ultima))
        aba.cell(linha, 4, f"=B{linha}-C{linha}")
        aba.cell(linha, 5, f"=D{linha}" if mes == 1 else f"=E{linha - 1}+D{linha}")
        for coluna in range(1, 6):
            celula = aba.cell(linha, coluna)
            celula.font = Font(name=FONTE, bold=coluna == 4)
            celula.alignment = Alignment(horizontal="center")
            borda(celula)
            if coluna > 1:
                celula.number_format = REAIS

    linha_total = 18
    aba.cell(linha_total, 1, "Total do ano")
    for coluna, letra in ((2, "B"), (3, "C"), (4, "D")):
        aba.cell(linha_total, coluna, f"=SUM({letra}6:{letra}17)").number_format = REAIS
    for coluna in range(1, 6):
        celula = aba.cell(linha_total, coluna)
        celula.font = Font(name=FONTE, bold=True)
        celula.fill = preencher(CINZA_CLARO)
        celula.alignment = Alignment(horizontal="center")

    # Saldo do mês: verde se sobrou dinheiro, vermelho se faltou
    aba.conditional_formatting.add("D6:D17", CellIsRule(
        operator="greaterThan", formula=["0"], fill=preencher(VERDE_CLARO), font=Font(color=VERDE)))
    aba.conditional_formatting.add("D6:D17", CellIsRule(
        operator="lessThan", formula=["0"], fill=preencher(VERMELHO_CLARO), font=Font(color=VERMELHO)))

    grafico = BarChart()
    grafico.title = "Receitas x Despesas"
    grafico.y_axis.title, grafico.y_axis.numFmt = "R$", '"R$" #,##0'
    grafico.add_data(Reference(aba, min_col=2, max_col=3, min_row=5, max_row=17), titles_from_data=True)
    grafico.set_categories(Reference(aba, min_col=1, min_row=6, max_row=17))
    grafico.series[0].graphicalProperties.solidFill = "2E9E5B"
    grafico.series[1].graphicalProperties.solidFill = "D64545"
    grafico.height, grafico.width = 9, 18
    aba.add_chart(grafico, "G5")


# ---------------------------------------------------------------------------
# Aba Orçamento: limite de cada categoria e o alerta de cores
# ---------------------------------------------------------------------------
def aba_orcamento(livro, ultima, mes_inicial):
    aba = livro.create_sheet("Orçamento")
    titulo(aba, "Orçamento do mês", "Verde: dentro do limite. Amarelo: passou de 80%. Vermelho: estourou.")
    aba["A3"] = "Mês:"
    aba["A3"].font = Font(name=FONTE, bold=True)
    aba["B3"] = mes_inicial
    aba["B3"].number_format = "mmm/yyyy"
    aba["B3"].font = Font(name=FONTE, bold=True, size=12, color=AZUL_ESCURO)
    aba["B3"].fill = preencher(AMARELO_CLARO)
    aba["C3"] = "← escolha o mês no menu"
    aba["C3"].font = Font(name=FONTE, italic=True, color="666666")
    meses = DataValidation(type="list", formula1="='Resumo mensal'!$A$6:$A$17", allow_blank=False)
    aba.add_data_validation(meses)
    meses.add("B3")

    cabecalho(aba, 5, [("Categoria", 32), ("Limite do mês", 15), ("Gasto", 15),
                       ("Ainda posso gastar", 19), ("Usado", 10), ("Situação", 14)])
    l = "'Lançamentos'!"
    primeira = 6
    for linha, (categoria, limite) in enumerate(DESPESAS.items(), start=primeira):
        aba.cell(linha, 1, categoria)
        aba.cell(linha, 2, limite)
        aba.cell(linha, 3, f'=SUMIFS({l}$E$5:$E${ultima},{l}$C$5:$C${ultima},"Despesa",'
                           f'{l}$D$5:$D${ultima},$A{linha},'
                           f'{l}$A$5:$A${ultima},">="&$B$3,{l}$A$5:$A${ultima},"<"&EDATE($B$3,1))')
        aba.cell(linha, 4, f"=B{linha}-C{linha}")
        aba.cell(linha, 5, f'=IF(B{linha}=0,0,C{linha}/B{linha})').number_format = "0%"
        aba.cell(linha, 6, f'=IF(C{linha}>B{linha},"Estourou",IF(C{linha}>=0.8*B{linha},"Atenção","Dentro"))')
        for coluna in (2, 3, 4):
            aba.cell(linha, coluna).number_format = REAIS
        for coluna in range(1, 7):
            celula = aba.cell(linha, coluna)
            celula.font = Font(name=FONTE)
            borda(celula)
            if coluna > 1:
                celula.alignment = Alignment(horizontal="center")
    final = primeira + len(DESPESAS) - 1

    total = final + 1
    aba.cell(total, 1, "Total")
    for coluna, letra in ((2, "B"), (3, "C"), (4, "D")):
        aba.cell(total, coluna, f"=SUM({letra}{primeira}:{letra}{final})").number_format = REAIS
    aba.cell(total, 5, f"=IF(B{total}=0,0,C{total}/B{total})").number_format = "0%"
    for coluna in range(1, 7):
        celula = aba.cell(total, coluna)
        celula.font = Font(name=FONTE, bold=True)
        celula.fill = preencher(CINZA_CLARO)
        celula.alignment = Alignment(horizontal="center" if coluna > 1 else None)

    # O alerta de cores: a linha inteira muda conforme a Situação
    faixa = f"A{primeira}:F{final}"
    for texto, fundo, letra in (("Estourou", VERMELHO_CLARO, VERMELHO),
                                ("Atenção", AMARELO_CLARO, AMARELO),
                                ("Dentro", VERDE_CLARO, VERDE)):
        aba.conditional_formatting.add(faixa, FormulaRule(
            formula=[f'$F{primeira}="{texto}"'], fill=preencher(fundo), font=Font(color=letra)))

    pizza = PieChart()
    pizza.title = "Para onde foi o dinheiro no mês"
    pizza.add_data(Reference(aba, min_col=3, min_row=5, max_row=final), titles_from_data=True)
    pizza.set_categories(Reference(aba, min_col=1, min_row=primeira, max_row=final))
    pizza.dataLabels = DataLabelList()           # mostra só a porcentagem em cada fatia
    pizza.dataLabels.showPercent = True
    for detalhe in ("showVal", "showCatName", "showSerName", "showLegendKey"):
        setattr(pizza.dataLabels, detalhe, False)
    pizza.legend.position = "r"
    pizza.height, pizza.width = 10, 16
    aba.add_chart(pizza, "H3")


# ---------------------------------------------------------------------------
# Aba Como usar
# ---------------------------------------------------------------------------
def aba_como_usar(aba):
    titulo(aba, "Minha planilha financeira", "Controle de receitas, despesas e orçamento, mês a mês.")
    passos = [
        ("1. Lançamentos", "Anote cada entrada e saída: data, descrição, tipo, categoria e valor. "
                           "Tipo e Categoria têm menu; o valor é sempre positivo."),
        ("2. Orçamento", "Escolha o mês no topo e veja quanto já gastou em cada categoria. "
                         "Verde = dentro do limite, amarelo = passou de 80%, vermelho = estourou."),
        ("3. Resumo mensal", "Receitas, despesas e saldo de cada mês, com gráfico. Mude o ano na célula amarela."),
        ("4. Categorias", "Crie ou renomeie categorias. Para ter limite no orçamento, "
                          "acrescente a categoria também na aba Orçamento."),
        ("Dica", "Células amarelas são para você mudar. As outras se calculam sozinhas."),
    ]
    aba.column_dimensions["A"].width = 20
    aba.column_dimensions["B"].width = 90
    for linha, (passo, explicacao) in enumerate(passos, start=4):
        aba.cell(linha, 1, passo).font = Font(name=FONTE, bold=True, color=AZUL_ESCURO)
        celula = aba.cell(linha, 2, explicacao)
        celula.font = Font(name=FONTE)
        celula.alignment = Alignment(wrap_text=True, vertical="top")
        aba.row_dimensions[linha].height = 32


# ---------------------------------------------------------------------------
# Lançamentos de exemplo: três meses de uma pessoa fictícia
# ---------------------------------------------------------------------------
def lancamentos_de_exemplo():
    sorteio = random.Random(2026)        # semente fixa: sempre os mesmos exemplos
    lancamentos = []
    for mes in (8, 9, 10):
        def dia(d):
            return date(ANO, mes, d)
        lancamentos += [
            (dia(5), "Salário", "Receita", "Salário", 4200.00),
            (dia(6), "Aluguel", "Despesa", "Moradia", 1450.00),
            (dia(10), "Conta de luz", "Despesa", "Contas (luz, água, internet)",
             round(sorteio.uniform(110, 170), 2)),
            (dia(10), "Internet", "Despesa", "Contas (luz, água, internet)", 99.90),
            (dia(12), "Água", "Despesa", "Contas (luz, água, internet)", round(sorteio.uniform(60, 90), 2)),
            (dia(15), "Curso online", "Despesa", "Educação", 79.90),
            (dia(20), "Farmácia", "Despesa", "Saúde", round(sorteio.uniform(40, 160), 2)),
        ]
        for semana in range(4):
            lancamentos.append((dia(3 + semana * 7), "Supermercado", "Despesa", "Mercado",
                                round(sorteio.uniform(170, 290), 2)))
            lancamentos.append((dia(4 + semana * 7), "Combustível / ônibus", "Despesa", "Transporte",
                                round(sorteio.uniform(70, 120), 2)))
        for _ in range(sorteio.randint(2, 4)):
            lancamentos.append((dia(sorteio.randint(1, 28)), sorteio.choice(
                ["Cinema", "Restaurante", "Show", "Passeio"]), "Despesa", "Lazer",
                round(sorteio.uniform(40, 150), 2)))
        if mes != 9:
            lancamentos.append((dia(18), "Projeto freela", "Receita", "Freela",
                                round(sorteio.uniform(300, 900), 2)))
        lancamentos.append((dia(25), "Presente de aniversário", "Despesa", "Outras despesas",
                            round(sorteio.uniform(60, 180), 2)))
    return sorted(lancamentos)


def main():
    if ARQUIVO.exists() and "--substituir" not in sys.argv:
        print(f"A planilha {ARQUIVO} já existe e não foi alterada (ela pode ter os seus lançamentos).")
        print("Para gerar outra por cima, rode: python gerar_planilha.py --substituir")
        return
    em_branco = "--em-branco" in sys.argv
    exemplos = [] if em_branco else lancamentos_de_exemplo()

    livro = Workbook()
    aba_como_usar(livro.active)
    livro.active.title = "Como usar"
    lista = aba_categorias(livro)
    ultima = aba_lancamentos(livro, lista, exemplos)
    mes_inicial = max(e[0] for e in exemplos).replace(day=1) if exemplos else date.today().replace(day=1)
    aba_orcamento(livro, ultima, mes_inicial)
    aba_resumo(livro, ultima)
    # Ordem das abas na tela: Como usar, Lançamentos, Orçamento, Resumo, Categorias
    livro._sheets = [livro[n] for n in ("Como usar", "Lançamentos", "Orçamento",
                                        "Resumo mensal", "Categorias")]
    livro.active = 1
    for aba in livro.worksheets:
        para_impressao(aba)
    # Imprime só as linhas preenchidas dos lançamentos, não as 1000 reservadas
    livro["Lançamentos"].print_area = f"A1:E{max(5, 4 + len(exemplos))}"

    try:
        livro.save(ARQUIVO)
    except PermissionError:
        print(f"Feche a planilha {ARQUIVO} no Excel e rode de novo.")
        sys.exit(1)
    print(f"Pronto: {ARQUIVO} ({len(exemplos)} lançamentos de exemplo).")


if __name__ == "__main__":
    main()
